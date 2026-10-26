"""Parse PDF/DOCX/PPTX into selectable text chunks."""
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
from pathlib import Path


def parse_pdf(path: str) -> list[dict]:
    doc = fitz.open(path)
    chunks = []
    for page_num, page in enumerate(doc, start=1):
        for block in page.get_text("blocks"):
            text = block[4].strip()
            if len(text) > 40:
                chunks.append({
                    "id": f"p{page_num}_b{len(chunks)}",
                    "text": text,
                    "page": page_num,
                    "type": "paragraph",
                })
    return chunks


def parse_docx(path: str) -> list[dict]:
    doc = Document(path)
    chunks = []
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if len(text) > 40:
            chunks.append({
                "id": f"para_{i}",
                "text": text,
                "page": None,
                "type": "paragraph",
            })
    return chunks


def parse_pptx(path: str) -> list[dict]:
    prs = Presentation(path)
    chunks = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text = shape.text.strip()
                if len(text) > 20:
                    chunks.append({
                        "id": f"s{slide_num}_t{len(chunks)}",
                        "text": text,
                        "page": slide_num,
                        "type": "paragraph",
                    })
    return chunks


def parse_document(path: str) -> list[dict]:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path)
    elif ext == ".docx":
        return parse_docx(path)
    elif ext == ".pptx":
        return parse_pptx(path)
    raise ValueError(f"Unsupported format: {ext}")
